"""All names, dates and values here are synthetic. Run from backend."""
from pathlib import Path
from datetime import datetime
import csv
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

ROOT = Path(__file__).parent
STANDARD = [
    ["Date", "Description", "Amount", "Category", "Type"],
    ["2026-08-01", "Synthetic salary", "1000.00", "Salary", "income"],
    ["2026-08-02", "Synthetic groceries", "-200.00", "Groceries", "expense"],
    ["2026-08-03", "Synthetic cafe", "-50.00", "Dining", "expense"],
    ["2026-08-04", "Synthetic savings movement", "-100.00", "", "transfer"],
]


def generate():
    variants = {
        "standard.csv": STANDARD,
        "debit_credit.csv": [["Date", "Reference", "Money Out", "Money In"],
                             ["01/08/2026", "Synthetic salary", "", "2000.00"],
                             ["02/08/2026", "Synthetic rent", "500.00", ""],
                             ["03/08/2026", "Synthetic fuel", "125.00", ""]],
        "malformed.csv": [["Date", "Description", "Amount"], ["31/02/2026", "Impossible date", "-50"],
                          ["2026-08-02", "Invalid amount", "not a number"]],
        "missing_column.csv": [["Reference", "Value"], ["Synthetic cafe", "-50.00"]],
        "duplicates.csv": STANDARD + [STANDARD[2]],
        "positive_expense.csv": [["Date", "Description", "Amount"],
                                ["2026-08-01", "Synthetic salary", "-1500.00"], ["2026-08-02", "Synthetic cafe", "123.45"]],
        "manual.csv": [["When", "Who", "Net"], ["2026-08-01", "Synthetic client", "800.10"]],
    }
    for name, rows in variants.items():
        with (ROOT / name).open("w", newline="", encoding="utf-8") as file:
            csv.writer(file).writerows(rows)
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Notes"
    notes.append(["Synthetic fixture only. No real bank data."])
    sheet = workbook.create_sheet("Transactions")
    sheet.append(["SYNTHETIC ACCOUNT EXPORT"])
    sheet.append(["An irrelevant title row"])
    sheet.append(["Transaction Date", "Narrative", "Value", "Running Balance", "Type", "Category", "Ignore me"])
    for row in [
        [datetime(2026, 8, 1), "Synthetic salary", 3000, 3000, "income", "Salary", "x"],
        [datetime(2026, 8, 2), "Synthetic groceries", -600, 2400, "expense", "Groceries", "x"],
        [datetime(2026, 8, 3), "Synthetic phone", -99.99, 2300.01, "expense", "Phone", "x"],
        [datetime(2026, 8, 4), "Synthetic transfer", -250, 2050.01, "transfer", "", "x"],
    ]:
        sheet.append(row)
    workbook.save(ROOT / "renamed_headers.xlsx")
    simple = Workbook()
    for row in STANDARD:
        simple.active.append(row)
    simple.save(ROOT / "standard.xlsx")
    import xlwt
    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet("Transactions")
    for i, row in enumerate(STANDARD):
        for j, value in enumerate(row):
            sheet.write(i, j, value)
    legacy.save(str(ROOT / "standard.xls"))
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table as OTable, TableRow, TableCell
    from odf.text import P
    ods = OpenDocumentSpreadsheet()
    sheet = OTable(name="Transactions")
    for row in STANDARD:
        tr = TableRow()
        for value in row:
            cell = TableCell(valuetype="string")
            cell.addElement(P(text=value))
            tr.addElement(cell)
        sheet.addElement(tr)
    ods.spreadsheet.addElement(sheet)
    ods.save(str(ROOT / "standard.ods"))
    table = Table(STANDARD, colWidths=[70, 175, 65, 65, 60], repeatRows=1)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                               ("FONTSIZE", (0, 0), (-1, -1), 8),
                               ("TOPPADDING", (0, 0), (-1, -1), 10),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    SimpleDocTemplate(str(ROOT / "synthetic_statement.pdf")).build([
        Paragraph("FinTrack synthetic statement — no real financial data", getSampleStyleSheet()["Heading2"]), table])


if __name__ == "__main__":
    generate()
